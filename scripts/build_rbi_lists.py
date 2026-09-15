"""Converts the official RBI Digital Lending Apps (DLA) directory export
(an .xlsx downloaded by hand from https://data.rbi.org.in/BOE/OpenDocument/
opendoc/custom.jsp?sIDType=CUID&iDocID=ARfEgy.WNSVIvFfvSIVmBCw — RBI's site
blocks automated/non-browser fetches, so this file must be exported by a
human in a real browser session each time it's refreshed) into the two JSON
files the project actually reads:

- data/rbi-regulated-lending-apps.json   (Node backend's GET /api/lending-check)
- backend-api/data/rbi_apps.json         (Python worker's rbi_checker.py)

Both are generated from this one source so they can't silently drift apart.

Usage: python3 scripts/build_rbi_lists.py path/to/DigitalLendingApp.xlsx
"""

import json
import sys
from datetime import date, datetime
from pathlib import Path

import openpyxl

REPO_ROOT = Path(__file__).resolve().parent.parent
NODE_OUTPUT = REPO_ROOT / "data" / "rbi-regulated-lending-apps.json"
PYTHON_OUTPUT = REPO_ROOT / "backend-api" / "data" / "rbi_apps.json"

SOURCE_URL = (
    "https://data.rbi.org.in/BOE/OpenDocument/opendoc/custom.jsp"
    "?sIDType=CUID&iDocID=ARfEgy.WNSVIvFfvSIVmBCw"
)
SOURCE_LABEL = (
    "Reserve Bank of India — Digital Lending Apps (DLA) directory "
    f"({SOURCE_URL}), via Citizen's Corner > \"DLA's deployed by "
    "Regulated Entities\""
)

# Column layout of the export, 0-indexed (row = tuple of cell values A.. M):
# 1 Sr.No | 2 Entity Name | 3 Website of RE | 4 Entity Type | 5 Name of the DLA
# | 6 Owner of DLA | 7 Available on | 8 Link to DLA | 9 Grievance officer name
# | 10 email | 11 phone | 12 mobile
COL_ENTITY_NAME = 2
COL_ENTITY_TYPE = 4
COL_DLA_NAME = 5
COL_LINK = 8


# A handful of rows in RBI's own export have a bare digit ("0".."5") or
# "NA" as the DLA name — data-entry errors from the self-reporting entity,
# not real app names. Matching against these would produce nonsense
# results (e.g. typing "3" would "match" a regulated app), so they're
# dropped. Nothing else about the source data is altered or guessed at.
_JUNK_NAMES = {"na", "n/a", "0", "1", "2", "3", "4", "5", "6", "7", "8", "9"}


# RBI's export merges the Entity Name cell down across all of that
# entity's app rows (a normal Excel "grouped rows" layout) — openpyxl
# returns None for every merged cell except the top-left one, so without
# forward-filling from the merge range, every app row after an entity's
# first would silently lose its entity attribution.
def _build_merged_entity_lookup(ws) -> dict:
    lookup = {}
    for rng in ws.merged_cells.ranges:
        if rng.min_col != COL_ENTITY_NAME + 1:  # openpyxl cols are 1-indexed
            continue
        top_value = ws.cell(row=rng.min_row, column=rng.min_col).value
        for r in range(rng.min_row, rng.max_row + 1):
            lookup[r] = top_value
    return lookup


def load_rows(xlsx_path: Path):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb.active
    merged_entities = _build_merged_entity_lookup(ws)
    rows = []
    dropped = []
    for row_idx, row in enumerate(ws.iter_rows(min_row=5, values_only=True), start=5):
        entity = row[COL_ENTITY_NAME]
        if entity is None:
            entity = merged_entities.get(row_idx)
        dla_name = row[COL_DLA_NAME]
        if entity is None and dla_name is None:
            continue
        if not dla_name or not str(dla_name).strip():
            continue  # can't match on an app with no name
        if str(dla_name).strip().lower() in _JUNK_NAMES:
            dropped.append(str(dla_name).strip())
            continue
        rows.append(
            {
                "name": str(dla_name).strip(),
                "entity": str(entity).strip() if entity else None,
                "entityType": str(row[COL_ENTITY_TYPE]).strip() if row[COL_ENTITY_TYPE] else None,
                "link": str(row[COL_LINK]).strip() if row[COL_LINK] else None,
            }
        )
    if dropped:
        print(f"Dropped {len(dropped)} junk/placeholder DLA names: {dropped}")
    return rows


def write_node_json(rows, today: str):
    seen = set()
    apps = []
    for r in rows:
        key = r["name"].lower()
        if key in seen:
            continue
        seen.add(key)
        apps.append(r["name"])
    apps.sort(key=str.lower)

    NODE_OUTPUT.write_text(
        json.dumps({"source": SOURCE_LABEL, "lastUpdated": today, "apps": apps}, indent=2)
        + "\n",
        encoding="utf-8",
    )
    print(f"Wrote {NODE_OUTPUT} ({len(apps)} unique app names)")


def write_python_json(rows, today: str):
    entries = [
        {
            "name": r["name"],
            "aliases": [],
            "entity": r["entity"],
            "entityType": r["entityType"],
            "link": r["link"],
        }
        for r in rows
    ]
    PYTHON_OUTPUT.write_text(
        json.dumps(
            {"source": SOURCE_LABEL, "last_updated": today, "apps": entries},
            indent=2,
        )
        + "\n",
        encoding="utf-8",
    )
    print(f"Wrote {PYTHON_OUTPUT} ({len(entries)} entries, one per RE/app row)")


def main():
    if len(sys.argv) != 2:
        print(__doc__)
        sys.exit(1)
    xlsx_path = Path(sys.argv[1]).expanduser().resolve()
    if not xlsx_path.exists():
        print(f"File not found: {xlsx_path}")
        sys.exit(1)

    rows = load_rows(xlsx_path)
    today = date.today().isoformat()
    write_node_json(rows, today)
    write_python_json(rows, today)


if __name__ == "__main__":
    main()
